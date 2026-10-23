"""Parser for multi-row Topia and Equus Excel workbooks."""

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl


@dataclass
class ColumnMetadata:
    col_index: int
    rule: str | None = None
    technical_field: str | None = None
    field_name: str | None = None
    display_name: str | None = None


@dataclass
class ParsedSheet:
    sheet_name: str
    columns: list[ColumnMetadata] = field(default_factory=list)
    data_rows: list[dict[int, Any]] = field(default_factory=list)
    header_rows_count: int = 3
    total_rows_evaluated: int = 0
    warnings: list[dict[str, Any]] = field(default_factory=list)


def parse_three_row_workbook(
    file_path: Path | str, sheet_name: str | None = None, max_rows: int | None = None
) -> ParsedSheet:
    """Parse a workbook with multi-row headers (Topia / Equus format).

    Row 1: Rules & transformation instructions (e.g. 'US Date Format', 'Match dropdown')
    Row 2: Source technical field identifier (e.g. 'EMPLOYEE!FIRST_NAME')
    Row 3: Destination import header / field name (e.g. 'First Name')
    Row 4: Optional display header (e.g. 'Employee Full Name') or first data row.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Workbook not found: {path}")

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    try:
        target_name = sheet_name or wb.sheetnames[0]
        if target_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{target_name}' not found in workbook")
        ws = wb[target_name]

        all_rows: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True):
            all_rows.append(list(row))
            if max_rows and len(all_rows) >= max_rows + 5:
                break
    finally:
        wb.close()

    if not all_rows:
        return ParsedSheet(sheet_name=target_name or "Sheet1")

    # Detect header depth: 3 vs 4 rows
    header_depth = 3
    if len(all_rows) >= 5:
        r4 = all_rows[3]
        r4_c1 = str(r4[0]).strip().upper() if len(r4) > 0 and r4[0] is not None else ""
        if r4_c1 in {"INSERT", "UPDATE"}:
            header_depth = 3
        else:
            # Check if row 4 contains header column titles
            r4_titles = [str(v).strip().lower() for v in r4[:10] if v is not None]
            if any(
                t
                in {
                    "employee full name",
                    "employee id",
                    "first name",
                    "last name",
                    "talent link id",
                    "email",
                }
                for t in r4_titles
            ):
                header_depth = 4

    max_cols = max(len(r) for r in all_rows)
    r1 = all_rows[0] if len(all_rows) > 0 else []
    r2 = all_rows[1] if len(all_rows) > 1 else []
    r3 = all_rows[2] if len(all_rows) > 2 else []
    r4 = all_rows[3] if len(all_rows) > 3 and header_depth == 4 else []

    columns: list[ColumnMetadata] = []
    seen_tech_fields: set[str] = set()
    warnings: list[dict[str, Any]] = []

    for c_idx in range(max_cols):
        c1 = str(r1[c_idx]).strip() if c_idx < len(r1) and r1[c_idx] is not None else None
        c2 = str(r2[c_idx]).strip() if c_idx < len(r2) and r2[c_idx] is not None else None
        c3 = str(r3[c_idx]).strip() if c_idx < len(r3) and r3[c_idx] is not None else None
        c4 = str(r4[c_idx]).strip() if c_idx < len(r4) and r4[c_idx] is not None else None

        # Clean instruction banners like 'Instructions --->>>'
        if c1 and "--->" in c1:
            c1 = None
        if c2 and "--->" in c2:
            c2 = None
        if c3 and "--->" in c3:
            c3 = None

        if c2:
            if c2 in seen_tech_fields:
                warnings.append(
                    {
                        "code": "DUPLICATE_TECHNICAL_FIELD",
                        "field": c2,
                        "column_index": c_idx + 1,
                        "message": (
                            f"Duplicate technical field identifier '{c2}' at column {c_idx + 1}"
                        ),
                    }
                )

            seen_tech_fields.add(c2)

        columns.append(
            ColumnMetadata(
                col_index=c_idx + 1,
                rule=c1,
                technical_field=c2,
                field_name=c3,
                display_name=c4,
            )
        )

    # Parse data rows starting after header rows
    data_rows: list[dict[int, Any]] = []
    data_start = header_depth
    total_evaluated = 0

    for r_idx in range(data_start, len(all_rows)):
        raw_row = all_rows[r_idx]
        total_evaluated += 1
        # Skip fully empty rows
        if not any(v is not None and str(v).strip() != "" for v in raw_row):
            continue

        row_dict: dict[int, Any] = {}
        for c_idx, val in enumerate(raw_row):
            if val is not None:
                # Strip string values
                if isinstance(val, str):
                    val = val.strip()
                row_dict[c_idx + 1] = val
        data_rows.append(row_dict)

    return ParsedSheet(
        sheet_name=target_name,
        columns=columns,
        data_rows=data_rows,
        header_rows_count=header_depth,
        total_rows_evaluated=total_evaluated,
        warnings=warnings,
    )
