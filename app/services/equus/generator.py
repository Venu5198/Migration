"""Artifact generation for Equus.

Produces the import Excel file and mapping/reconciliation report.
"""

from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.equus.transformer import TransformationResult


def generate_equus_import_file(
    template_path: Path | str,
    records: list[dict[int, Any]],
    output_path: Path | str,
    overwrite: bool = False,
) -> Path:
    """Generate the Equus import Excel workbook using the bundled template.

    Preserves row 1-3 formatting, headers, and rules.
    Writes one INSERT record per assignment + UPDATE records for child entities.
    """
    in_path = Path(template_path)
    out_path = Path(output_path)
    if out_path.exists() and not overwrite:
        raise FileExistsError(
            f"Target import file '{out_path}' already exists; will not overwrite."
        )
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(in_path)
    ws = wb.active

    # Clear any sample rows starting at row 4
    if ws.max_row >= 4:
        ws.delete_rows(4, ws.max_row - 3)

    # Write each transformed record into row 4, 5, ...
    for r_idx, record in enumerate(records, start=4):
        for col_idx, val in record.items():
            if val is not None:
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.value = val

    wb.save(out_path)
    wb.close()
    return out_path


def generate_mapping_report_file(
    result: TransformationResult,
    source_name: str,
    output_path: Path | str,
    exec_summary: Any = None,
    overwrite: bool = False,
    recon_result: dict[str, Any] | None = None,
) -> Path:
    """Generate a comprehensive multi-tab Excel mapping and reconciliation report."""
    out_path = Path(output_path)
    if out_path.exists() and not overwrite:
        raise FileExistsError(
            f"Target report file '{out_path}' already exists; will not overwrite."
        )
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=10)
    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # Tab 1: Executive Summary
    ws_summary = wb.create_sheet(title="Executive Summary")
    ws_summary.views.sheetView[0].showGridLines = True

    recon_status = (
        recon_result.get("status")
        if recon_result and "status" in recon_result
        else ("VERIFIED_PASS" if not result.warnings else "COMPLETED_WITH_WARNINGS")
    )
    val_errors = recon_result.get("errors", []) if recon_result else []
    strategy_used = (
        recon_result.get("assignment_id_strategy")
        if recon_result and "assignment_id_strategy" in recon_result
        else getattr(result, "strategy", "EMPLOYEE_NUMBER")
    )
    accepted_records = (
        recon_result.get("accepted_records")
        if recon_result and "accepted_records" in recon_result
        else (len(result.all_output_records) if recon_status == "VERIFIED_PASS" else 0)
    )
    rejected_records = (
        recon_result.get("rejected_records")
        if recon_result and "rejected_records" in recon_result
        else (0 if recon_status == "VERIFIED_PASS" else len(result.all_output_records))
    )

    summary_data = [
        ("Equus Migration & Reconciliation Summary", ""),
        ("Source Workbook", source_name),
        ("Execution Timestamp (UTC)", datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S")),
        ("Assignment Identifier Strategy", strategy_used),
        ("Total Source Rows Evaluated", result.total_evaluated_rows),
        ("Generated INSERT Records (Assignments)", len(result.insert_records)),
        ("Generated UPDATE Records (Child Entities)", len(result.update_records)),
        ("Total Records in Equus Import File", len(result.all_output_records)),
        ("Accepted Import Records", accepted_records),
        ("Rejected / Diagnostic Records", rejected_records),
        ("Equus Template Target Columns", len(result.mappings)),
        ("Mapped Target Columns", result.mapped_columns_count),
        ("Unmapped Target Columns", result.unmapped_columns_count),
        ("Field Coverage (%)", f"{result.coverage_percentage}%"),
        ("Validation Warnings", len(result.warnings)),
        ("Validation Errors", len(val_errors)),
        ("Dropdown Translations Applied", len(result.dropdown_translations)),
        ("Reconciliation Status", recon_status),
    ]


    if exec_summary and getattr(exec_summary, "demo_writes_enabled", False):
        summary_data.extend(
            [
                ("--- Demo Destination Execution ---", ""),
                ("Demo Target Destination", getattr(exec_summary, "target", "mock-equus-local")),
                ("Demo Operations Attempted", getattr(exec_summary, "attempted", 0)),
                ("Demo Operations Successful", getattr(exec_summary, "successful", 0)),
                ("Records Inserted", getattr(exec_summary, "inserted", 0)),
                ("Records Updated", getattr(exec_summary, "updated", 0)),
                ("Records Skipped (Idempotent)", getattr(exec_summary, "skipped", 0)),
                ("Records Failed", getattr(exec_summary, "failed", 0)),
                (
                    "Destination Read-Back Verification",
                    getattr(exec_summary, "verification_status", "UNVERIFIED"),
                ),
            ]
        )
    else:
        summary_data.extend(
            [
                ("--- Target Execution Status ---", ""),
                ("Destination Execution Status", "PLANNED ONLY (Target writes disabled by policy)"),
                ("Target Writes Attempted", 0),
                ("Target Writes Completed", 0),
            ]
        )
    for row_idx, (k, v) in enumerate(summary_data, start=1):
        c1 = ws_summary.cell(row=row_idx, column=1, value=k)
        c2 = ws_summary.cell(row=row_idx, column=2, value=v)
        c1.font = bold_font
        c2.font = regular_font
        if row_idx == 1:
            c1.font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
            ws_summary.merge_cells("A1:B1")

    # Tab 2: Column Mappings
    ws_map = wb.create_sheet(title="Column Mappings")
    ws_map.views.sheetView[0].showGridLines = True
    headers_map = [
        "Equus Col #",
        "Equus Technical Field (Row 3)",
        "Equus Field Label (Row 2)",
        "Equus Transformation Rule (Row 1)",
        "Mapping Kind",
        "Source Col #",
        "Source Technical Field",
        "Source Field Name",
        "Source Display Header",
    ]
    for col_idx, h in enumerate(headers_map, start=1):
        cell = ws_map.cell(row=1, column=col_idx, value=h)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, m in enumerate(result.mappings, start=2):
        row_vals = [
            m.eq_col_idx,
            m.eq_header or "",
            m.eq_label or "",
            m.eq_rule or "",
            m.mapping_kind,
            m.src_col_idx if m.src_col_idx else "-",
            m.src_tech_field or "-",
            m.src_field_name or "-",
            m.src_display_name or "-",
        ]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws_map.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = border_thin

    # Tab 3: Dropdown & Rule Translations
    ws_trans = wb.create_sheet(title="Translations Applied")
    ws_trans.views.sheetView[0].showGridLines = True
    headers_trans = ["Source Row #", "Equus Target Field", "Translation Detail"]
    for col_idx, h in enumerate(headers_trans, start=1):
        cell = ws_trans.cell(row=1, column=col_idx, value=h)
        cell.fill = navy_fill
        cell.font = header_font

    for r_idx, t in enumerate(result.dropdown_translations, start=2):
        ws_trans.cell(row=r_idx, column=1, value=t.get("source_row")).font = regular_font
        ws_trans.cell(row=r_idx, column=2, value=t.get("field")).font = regular_font
        ws_trans.cell(row=r_idx, column=3, value=t.get("translation")).font = regular_font

    # Tab 4: Warnings & Data Issues
    ws_warn = wb.create_sheet(title="Warnings & Issues")
    ws_warn.views.sheetView[0].showGridLines = True
    headers_warn = ["Code", "Source Row #", "Equus Column #", "Equus Header", "Warning Message"]
    for col_idx, h in enumerate(headers_warn, start=1):
        cell = ws_warn.cell(row=1, column=col_idx, value=h)
        cell.fill = navy_fill
        cell.font = header_font

    # 1. Output reconciliation validation errors first (critical failures)
    if val_errors:
        import re
        for err in val_errors:
            row_num = ws_warn.max_row + 1
            err_str = str(err)
            code = (
                "REQUIRED_FIELD_ERROR"
                if "missing required" in err_str or "Col " in err_str
                else "RECONCILIATION_ERROR"
            )

            # Extract record/row number if present (e.g. "record #1")
            m_row = re.search(r"#(\d+)", err_str)
            src_row = int(m_row.group(1)) if m_row else ""

            # Extract column # if present (e.g. "Col 3", "Col 5")
            m_col = re.search(r"Col\s*(\d+)", err_str)
            eq_col = int(m_col.group(1)) if m_col else ""

            # Extract target header name
            eq_hdr = ""
            if "Employee Number" in err_str:
                eq_hdr = "Employee Number"
            elif "Assignment Identifier" in err_str:
                eq_hdr = "Assignment Identifier"
            elif "Assignment Number" in err_str:
                eq_hdr = "Assignment Number"
            elif "Assignment Command" in err_str:
                eq_hdr = "Assignment Command"

            c1 = ws_warn.cell(row=row_num, column=1, value=code)
            c2 = ws_warn.cell(row=row_num, column=2, value=src_row)
            c3 = ws_warn.cell(row=row_num, column=3, value=eq_col)
            c4 = ws_warn.cell(row=row_num, column=4, value=eq_hdr)
            c5 = ws_warn.cell(row=row_num, column=5, value=err_str)

            c1.font = bold_font
            c2.font = regular_font
            c3.font = regular_font
            c4.font = regular_font
            c5.font = regular_font

    # 2. Output transformation warnings
    for r_idx, w in enumerate(result.warnings, start=ws_warn.max_row + 1):
        ws_warn.cell(row=r_idx, column=1, value=w.get("code")).font = regular_font
        ws_warn.cell(row=r_idx, column=2, value=w.get("source_row")).font = regular_font
        ws_warn.cell(row=r_idx, column=3, value=w.get("equus_column")).font = regular_font
        ws_warn.cell(row=r_idx, column=4, value=w.get("equus_header")).font = regular_font
        ws_warn.cell(row=r_idx, column=5, value=w.get("warning")).font = regular_font


    if exec_summary and getattr(exec_summary, "errors", None):
        for err in exec_summary.errors:
            row_num = ws_warn.max_row + 1
            ws_warn.cell(row=row_num, column=1, value="DESTINATION_WRITE_ERROR").font = regular_font
            ws_warn.cell(row=row_num, column=2, value=err.get("record_index")).font = regular_font
            ws_warn.cell(row=row_num, column=5, value=err.get("error")).font = regular_font

    if exec_summary and getattr(exec_summary, "discrepancies", None):
        for disc in exec_summary.discrepancies:
            row_num = ws_warn.max_row + 1
            ws_warn.cell(row=row_num, column=1, value="READ_BACK_DISCREPANCY").font = regular_font
            ws_warn.cell(row=row_num, column=2, value=disc.get("record_index")).font = regular_font
            ws_warn.cell(row=row_num, column=4, value=disc.get("field")).font = regular_font
            ws_warn.cell(row=row_num, column=5, value=disc.get("error")).font = regular_font

    # Auto-adjust column widths for all sheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    wb.save(out_path)
    wb.close()
    return out_path
