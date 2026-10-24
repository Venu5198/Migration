import hashlib
import re
import uuid
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook

from app.core.errors import DomainError
from app.schemas import IssueSeverity, SheetSchema, ValidationIssue

ALLOWED_SUFFIXES = {".xlsx", ".xlsm"}
_SAFE_NAME = re.compile(r"[^A-Za-z0-9._-]+")


def safe_filename(filename: str) -> str:
    candidate = _SAFE_NAME.sub("_", Path(filename).name).strip("._")
    if not candidate:
        raise DomainError("INVALID_FILENAME", "The filename is not usable.")
    return candidate[:180]


def validate_excel_upload(
    filename: str, stream: BinaryIO, *, max_bytes: int, allow_xlsm: bool
) -> tuple[str, bytes]:
    cleaned = safe_filename(filename)
    suffix = Path(cleaned).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES or (suffix == ".xlsm" and not allow_xlsm):
        raise DomainError(
            "UNSUPPORTED_FILE_TYPE", "Only .xlsx files are accepted by current policy."
        )
    payload = stream.read(max_bytes + 1)
    if len(payload) > max_bytes:
        raise DomainError(
            "UPLOAD_TOO_LARGE", "The workbook exceeds the configured upload limit.", 413
        )
    if not payload.startswith(b"PK\x03\x04"):
        raise DomainError("INVALID_WORKBOOK", "The uploaded file is not an OOXML workbook.")
    return cleaned, payload


def save_workbook(payload: bytes, filename: str, upload_dir: Path) -> tuple[str, Path, str]:
    source_id = str(uuid.uuid4())
    upload_dir.mkdir(parents=True, exist_ok=True)
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        suffix = ".xlsx"
    target_path = (upload_dir / f"{source_id}{suffix}").resolve()
    resolved_dir = upload_dir.resolve()
    if not target_path.is_relative_to(resolved_dir):
        raise DomainError("PATH_TRAVERSAL_DETECTED", "Unsafe storage destination.")
    target_path.write_bytes(payload)
    return source_id, target_path, hashlib.sha256(payload).hexdigest()


def _type_name(value: object) -> str:
    if value is None:
        return "empty"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, (int, float)):
        return "number"
    return "string"


from dataclasses import dataclass, field


@dataclass
class SheetLayout:
    kind: str  # "multi_row", "single_row", "empty"
    header_rows_count: int = 1
    data_start_row: int = 2
    headers: list[str] = field(default_factory=list)
    technical_headers: list[str] = field(default_factory=list)
    column_lookup: dict[str, int] = field(default_factory=dict)


def detect_sheet_layout(ws) -> SheetLayout:
    """Analyze worksheet rows to determine single-row vs multi-row Topia/Equus layout."""
    all_rows: list[list[object]] = []
    for r in ws.iter_rows(values_only=True):
        all_rows.append(list(r))
        if len(all_rows) >= 10:
            break

    if not all_rows or not any(any(v not in (None, "") for v in row) for row in all_rows):
        return SheetLayout(kind="empty", header_rows_count=0, data_start_row=1)

    max_cols = max(len(r) for r in all_rows)
    r1 = all_rows[0] if len(all_rows) > 0 else []
    r2 = all_rows[1] if len(all_rows) > 1 else []
    r3 = all_rows[2] if len(all_rows) > 2 else []
    r4 = all_rows[3] if len(all_rows) > 3 else []

    # Check for multi-row Topia/Equus technical identifiers or rules
    has_tech = any("!" in str(v) for v in r2 if v is not None) or any(
        "!" in str(v) for v in r1 if v is not None
    )
    has_rules = any(
        any(k in str(v).lower() for k in ["date format", "dropdown", "instructions", "yes/no", "yesno"])
        for v in r1
        if v is not None
    )

    if has_tech or has_rules:
        header_depth = 3
        if len(all_rows) >= 5:
            r4_c1 = str(r4[0]).strip().upper() if len(r4) > 0 and r4[0] is not None else ""
            if r4_c1 in {"INSERT", "UPDATE"}:
                header_depth = 3
            else:
                r4_titles = [str(v).strip().lower() for v in r4 if v is not None]
                known_headers = {
                    "full name",
                    "employee full name",
                    "employee id",
                    "first name",
                    "last name",
                    "talent link id",
                    "email",
                    "birth date",
                    "assignment id",
                    "gender",
                    "home country",
                    "date of birth",
                }
                if any(t in known_headers for t in r4_titles):
                    header_depth = 4

        canonical_headers: list[str] = []
        technical_headers: list[str] = []
        column_lookup: dict[str, int] = {}

        for c in range(max_cols):
            c_tech = str(r2[c]).strip() if c < len(r2) and r2[c] is not None else ""
            c_label = str(r3[c]).strip() if c < len(r3) and r3[c] is not None else ""
            c_disp = (
                str(r4[c]).strip()
                if header_depth == 4 and c < len(r4) and r4[c] is not None
                else ""
            )

            # Clean banner text
            if "--->" in c_tech:
                c_tech = ""
            if "--->" in c_label:
                c_label = ""
            if "--->" in c_disp:
                c_disp = ""

            if c_disp:
                name = c_disp
            elif c_label:
                name = c_label
            elif c_tech:
                name = c_tech
            else:
                name = f"__blank_{c + 1}"

            canonical_headers.append(name)
            technical_headers.append(c_tech)

            # Register lookup aliases for mapping resolution
            for alias in [name, c_disp, c_label, c_tech]:
                if alias and not alias.startswith("__blank_"):
                    column_lookup[alias] = c
                    column_lookup[alias.lower()] = c

        return SheetLayout(
            kind="multi_row",
            header_rows_count=header_depth,
            data_start_row=header_depth + 1,
            headers=canonical_headers,
            technical_headers=technical_headers,
            column_lookup=column_lookup,
        )

    # Standard single-row worksheet
    canonical_headers = [
        str(v).strip() if v not in (None, "") else f"__blank_{i + 1}" for i, v in enumerate(r1)
    ]
    column_lookup = {}
    for i, h in enumerate(canonical_headers):
        if not h.startswith("__blank_"):
            column_lookup[h] = i
            column_lookup[h.lower()] = i

    return SheetLayout(
        kind="single_row",
        header_rows_count=1,
        data_start_row=2,
        headers=canonical_headers,
        technical_headers=[],
        column_lookup=column_lookup,
    )


def inspect_workbook(path: Path) -> list[SheetSchema]:
    try:
        workbook = load_workbook(
            path, read_only=True, data_only=False, keep_vba=path.suffix == ".xlsm"
        )
    except Exception as exc:
        raise DomainError("UNREADABLE_WORKBOOK", "The workbook could not be read.") from exc
    schemas: list[SheetSchema] = []
    for ws in workbook.worksheets:
        layout = detect_sheet_layout(ws)
        issues: list[ValidationIssue] = []

        if layout.kind == "empty":
            issues.append(
                ValidationIssue(
                    code="EMPTY_WORKSHEET",
                    severity=IssueSeverity.WARNING,
                    message="Worksheet has no detectable header.",
                    sheet=ws.title,
                )
            )
            schemas.append(
                SheetSchema(
                    name=ws.title,
                    max_row=ws.max_row or 0,
                    max_column=ws.max_column or 0,
                    issues=issues,
                )
            )
            continue

        headers = layout.headers
        duplicates = [header for header, count in Counter(headers).items() if count > 1]
        for header in duplicates:
            issues.append(
                ValidationIssue(
                    code="DUPLICATE_COLUMN",
                    severity=IssueSeverity.WARNING,
                    message="Duplicate column name requires mapping review.",
                    sheet=ws.title,
                    field=header,
                )
            )
        for header in [item for item in headers if item.startswith("__blank_")]:
            issues.append(
                ValidationIssue(
                    code="MISSING_HEADER",
                    severity=IssueSeverity.WARNING,
                    message="Blank header is preserved with a synthetic label.",
                    sheet=ws.title,
                    field=header,
                )
            )

        samples: dict[str, set[str]] = {header: set() for header in headers}
        rows = ws.iter_rows(values_only=True)
        # Fast forward past all header rows
        for _ in range(layout.data_start_row - 1):
            next(rows, None)

        for row_index, row in enumerate(rows, start=layout.data_start_row):
            if row_index > layout.data_start_row + 200:
                break
            for index, value in enumerate(row[: len(headers)]):
                samples[headers[index]].add(_type_name(value))

        inferred = {
            header: next(iter(values)) if len(values) == 1 else "mixed"
            for header, values in samples.items()
        }
        schemas.append(
            SheetSchema(
                name=ws.title,
                max_row=ws.max_row or 0,
                max_column=ws.max_column or 0,
                headers=headers,
                technical_headers=layout.technical_headers,
                header_rows_count=layout.header_rows_count,
                data_start_row=layout.data_start_row,
                inferred_types=inferred,
                issues=issues,
            )
        )
    workbook.close()
    return schemas

