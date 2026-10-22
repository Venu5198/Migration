from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from app.schemas import IssueSeverity, ValidationIssue


def profile_workbook(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    issues: list[dict] = []
    for ws in workbook.worksheets:
        values = ws.iter_rows(values_only=True)
        headers = next(values, None)
        if not headers:
            continue
        labels = [
            str(value).strip() if value not in (None, "") else f"__blank_{i + 1}"
            for i, value in enumerate(headers)
        ]
        nulls = Counter()
        populated_rows = 0
        structural_empty_rows = 0
        seen_rows = set()
        duplicates = 0
        for row_number, row in enumerate(values, start=2):
            normalized = tuple(row[: len(labels)])
            if not any(value not in (None, "") for value in normalized):
                structural_empty_rows += 1
                continue
            populated_rows += 1
            if normalized in seen_rows:
                duplicates += 1
                issues.append(
                    ValidationIssue(
                        code="DUPLICATE_ROW",
                        severity=IssueSeverity.WARNING,
                        message="Duplicate populated row detected.",
                        sheet=ws.title,
                        row=row_number,
                    ).model_dump()
                )
            seen_rows.add(normalized)
            for label, value in zip(labels, normalized, strict=False):
                if value in (None, ""):
                    nulls[label] += 1
        sheets.append(
            {
                "sheet": ws.title,
                "populated_rows": populated_rows,
                "empty_structural_rows": structural_empty_rows,
                "null_counts": dict(nulls),
                "duplicate_rows": duplicates,
            }
        )
    workbook.close()
    return {"sheets": sheets, "issues": issues}
