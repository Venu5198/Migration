from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app
from app.schemas import MappingType
from app.services.reporting import preview_csv
from app.services.transform import transform_value


@pytest.fixture
def client():
    with TestClient(app) as test_client:
        yield test_client


def workbook_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Assignments"
    worksheet.append(["Employee ID", "Country", "Country"])
    worksheet.append(["001", "India", "India"])
    worksheet.append([None, None, None])
    workbook.create_sheet("Empty")
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def upload(client: TestClient) -> dict:
    response = client.post(
        "/api/v1/sources/upload",
        files={
            "file": (
                "client.xlsx",
                workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_health_is_explicit_about_equus(client: TestClient):
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["equus_writes"] == "disabled"


def test_liveness_and_readiness(client: TestClient):
    assert client.get("/live").status_code == 200
    assert client.get("/ready").status_code == 200


def test_upload_discovers_all_sheets_and_flags_duplicate_header(client: TestClient):
    source = upload(client)
    assert {sheet["name"] for sheet in source["sheets"]} == {"Assignments", "Empty"}
    assignment = next(sheet for sheet in source["sheets"] if sheet["name"] == "Assignments")
    assert any(issue["code"] == "DUPLICATE_COLUMN" for issue in assignment["issues"])
    empty_sheet = next(sheet for sheet in source["sheets"] if sheet["name"] == "Empty")
    assert isinstance(empty_sheet["max_row"], int)
    assert isinstance(empty_sheet["max_column"], int)
    assert source["sha256"]


def test_inspect_workbook_handles_none_dimensions(tmp_path):
    from unittest.mock import MagicMock, patch

    from app.services.workbooks import inspect_workbook

    mock_wb = MagicMock()
    mock_ws = MagicMock()
    mock_ws.title = "EmptySheet"
    mock_ws.max_row = None
    mock_ws.max_column = None
    mock_ws.iter_rows.return_value = iter([])
    mock_wb.worksheets = [mock_ws]

    with patch("app.services.workbooks.load_workbook", return_value=mock_wb):
        schemas = inspect_workbook(tmp_path / "dummy.xlsx")
        assert len(schemas) == 1
        assert schemas[0].max_row == 0
        assert schemas[0].max_column == 0


def test_rejects_non_workbook_upload(client: TestClient):
    response = client.post(
        "/api/v1/sources/upload",
        files={"file": ("bad.xlsx", b"not a zip", "application/octet-stream")},
    )
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "INVALID_WORKBOOK"


def test_mapping_cannot_approve_unresolved(client: TestClient):
    source = upload(client)
    payload = {
        "source_id": source["source_id"],
        "source_entity": "employee",
        "source_sheet": "Assignments",
        "source_field": "Host Level",
        "target_entity": "assignment",
        "target_field": "host_job_title",
    }
    created = client.post("/api/v1/mappings", json=payload)
    assert created.status_code == 201
    denied = client.post(f"/api/v1/mappings/{created.json()['mapping_id']}/approve")
    assert denied.status_code == 409
    approved_payload = {
        **payload,
        "source_field": "Employee ID",
        "target_field": "employee_id",
        "mapping_type": MappingType.DIRECT,
    }
    approved = client.post("/api/v1/mappings", json=approved_payload)
    assert (
        client.post(f"/api/v1/mappings/{approved.json()['mapping_id']}/approve").status_code == 200
    )


@pytest.mark.asyncio
async def test_profile_job_is_claimed_and_completed(client: TestClient):
    from app.workers.worker import process_once

    source = upload(client)
    created = client.post(f"/api/v1/sources/{source['source_id']}/profile")
    assert created.status_code == 202
    assert await process_once(app.state.repository, "test-worker")
    job = client.get(f"/api/v1/jobs/{created.json()['job_id']}").json()
    assert job["status"] == "COMPLETED"
    assert job["result"]["sheets"][0]["populated_rows"] == 1


@pytest.mark.asyncio
async def test_direct_mapping_preview_never_writes_equus(client: TestClient):
    from app.workers.worker import process_once

    source = upload(client)
    created = client.post(
        "/api/v1/mappings",
        json={
            "source_id": source["source_id"],
            "source_entity": "employee",
            "source_sheet": "Assignments",
            "source_field": "Employee ID",
            "target_entity": "employee",
            "target_field": "employee_id",
            "mapping_type": "direct",
        },
    ).json()
    assert client.post(f"/api/v1/mappings/{created['mapping_id']}/approve").status_code == 200
    preview = client.post(
        "/api/v1/migrations/preview",
        json={"source_id": source["source_id"], "mapping_id": created["mapping_id"]},
    )
    assert preview.status_code == 202
    await process_once(app.state.repository, "test-worker")
    job = client.get(f"/api/v1/jobs/{preview.json()['job_id']}").json()
    assert job["status"] == "COMPLETED"
    assert job["result"]["records"][0]["proposed"]["employee_id"] == "001"
    assert job["result"]["equus_writes"] is False


def test_transform_is_deterministic_and_csv_escapes_formulas():
    assert transform_value("  india  ", {"op": "upper"}) == "  INDIA  "
    report = preview_csv(
        {
            "records": [
                {"sheet": "A", "row": 2, "target_entity": "employee", "proposed": {"name": "=1+1"}}
            ]
        }
    )
    assert "'=1+1" in report


def test_authentication_and_approval_role(client: TestClient, monkeypatch):
    from app.core.config import settings

    monkeypatch.setattr(settings, "auth_required", True)
    monkeypatch.setattr(
        settings, "auth_tokens", "reader:r:viewer,approver:a:approver,admin:z:admin"
    )
    assert client.get("/api/v1/sources").status_code == 401
    assert (
        client.get("/api/v1/sources", headers={"Authorization": "Bearer reader"}).status_code == 200
    )
    assert (
        client.get("/api/v1/audit-events", headers={"Authorization": "Bearer reader"}).status_code
        == 403
    )
    assert (
        client.get("/api/v1/audit-events", headers={"Authorization": "Bearer admin"}).status_code
        == 200
    )
    monkeypatch.setattr(settings, "auth_required", False)


def test_pagination_and_audit_lineage(client: TestClient):
    source = upload(client)
    page = client.get("/api/v1/sources?page=1&page_size=1")
    assert page.status_code == 200
    assert page.json()["total"] >= 1
    assert client.get("/api/v1/sources?page=0").status_code == 422
    assert client.get(f"/api/v1/sources/{source['source_id']}").status_code == 200
    actions = [event["action"] for event in app.state.repository.audits]
    assert "source.uploaded" in actions
    assert "source.accessed" in actions


@pytest.mark.asyncio
async def test_cancelled_job_is_never_claimed(client: TestClient):
    from app.workers.worker import process_once

    source = upload(client)
    job = client.post(f"/api/v1/sources/{source['source_id']}/profile").json()
    assert client.post(f"/api/v1/jobs/{job['job_id']}/cancel").status_code == 200
    assert not await process_once(app.state.repository, "test-worker")
    assert client.get(f"/api/v1/jobs/{job['job_id']}").json()["status"] == "CANCELLED"


def test_mapping_version_history(client: TestClient):
    source = upload(client)
    mapping = client.post(
        "/api/v1/mappings",
        json={
            "source_id": source["source_id"],
            "source_entity": "employee",
            "source_sheet": "Assignments",
            "source_field": "Employee ID",
            "target_entity": "employee",
            "target_field": "employee_id",
            "mapping_type": "direct",
        },
    ).json()
    assert client.post(f"/api/v1/mappings/{mapping['mapping_id']}/approve").status_code == 200
    assert (
        client.get(f"/api/v1/mappings/{mapping['mapping_id']}/versions").json()[0]["version"] == 2
    )


def test_production_persistence_fail_closed():
    from app.core.config import Settings

    prod_bad = Settings(env="production", mongodb_uri=None, auth_tokens="t:s:admin")
    with pytest.raises(ValueError, match="MIGRATION_MONGODB_URI is required in production"):
        prod_bad.validate_production()


def test_hardened_csv_formula_injection():
    from app.services.reporting import safe_csv

    assert safe_csv("  =1+1").startswith("'")
    assert safe_csv("\t+CMD()").startswith("'")
    assert safe_csv("   @SUM(A1)").startswith("'")
    assert safe_csv("   -5+5").startswith("'")
    assert safe_csv("normal text") == "normal text"


def test_mapping_reapproval_rejected(client: TestClient):
    source = upload(client)
    mapping = client.post(
        "/api/v1/mappings",
        json={
            "source_id": source["source_id"],
            "source_entity": "employee",
            "source_sheet": "Assignments",
            "source_field": "Employee ID",
            "target_entity": "employee",
            "target_field": "employee_id",
            "mapping_type": "direct",
        },
    ).json()
    assert client.post(f"/api/v1/mappings/{mapping['mapping_id']}/approve").status_code == 200
    reapprove = client.post(f"/api/v1/mappings/{mapping['mapping_id']}/approve")
    assert reapprove.status_code == 409
    assert reapprove.json()["error"]["code"] == "MAPPING_ALREADY_APPROVED"


def test_mapping_separation_of_duties(client: TestClient, monkeypatch):
    from app.core.config import settings

    source = upload(client)
    monkeypatch.setattr(settings, "auth_required", True)
    monkeypatch.setattr(
        settings,
        "auth_tokens",
        "token_author:author_user:executor|approver,token_reviewer:reviewer_user:approver",
    )
    created = client.post(
        "/api/v1/mappings",
        json={
            "source_id": source["source_id"],
            "source_entity": "employee",
            "source_sheet": "Assignments",
            "source_field": "Employee ID",
            "target_entity": "employee",
            "target_field": "employee_id",
            "mapping_type": "direct",
        },
        headers={"Authorization": "Bearer token_author"},
    ).json()
    self_approve = client.post(
        f"/api/v1/mappings/{created['mapping_id']}/approve",
        headers={"Authorization": "Bearer token_author"},
    )
    assert self_approve.status_code == 403
    assert self_approve.json()["error"]["code"] == "SEPARATION_OF_DUTIES_VIOLATION"

    reviewer_approve = client.post(
        f"/api/v1/mappings/{created['mapping_id']}/approve",
        headers={"Authorization": "Bearer token_reviewer"},
    )
    assert reviewer_approve.status_code == 200
    monkeypatch.setattr(settings, "auth_required", False)


def test_mapping_update_endpoint(client: TestClient):
    source = upload(client)
    mapping = client.post(
        "/api/v1/mappings",
        json={
            "source_id": source["source_id"],
            "source_entity": "employee",
            "source_sheet": "Assignments",
            "source_field": "Employee ID",
            "target_entity": "employee",
            "target_field": "employee_id",
            "mapping_type": "unresolved",
        },
    ).json()
    updated = client.put(
        f"/api/v1/mappings/{mapping['mapping_id']}",
        json={"mapping_type": "direct", "required": True},
    )
    assert updated.status_code == 200
    assert updated.json()["mapping_type"] == "direct"
    assert updated.json()["required"] is True

    assert client.post(f"/api/v1/mappings/{mapping['mapping_id']}/approve").status_code == 200

    attempt = client.put(
        f"/api/v1/mappings/{mapping['mapping_id']}",
        json={"required": False},
    )
    assert attempt.status_code == 409
    assert attempt.json()["error"]["code"] == "MAPPING_ALREADY_APPROVED"


@pytest.mark.asyncio
async def test_job_claim_exceeds_max_attempts(client: TestClient):
    repo = app.state.repository
    source = upload(client)
    created = client.post(f"/api/v1/sources/{source['source_id']}/profile").json()
    job_id = created["job_id"]
    job = await repo.get_job(job_id)
    job["attempts"] = 3
    await repo.put_job(job)
    claimed = await repo.claim_job("test-worker", max_attempts=3)
    assert claimed is None


@pytest.mark.asyncio
async def test_export_and_reconciliation_audit_events(client: TestClient):
    from app.workers.worker import process_once

    source = upload(client)
    created = client.post(
        "/api/v1/mappings",
        json={
            "source_id": source["source_id"],
            "source_entity": "employee",
            "source_sheet": "Assignments",
            "source_field": "Employee ID",
            "target_entity": "employee",
            "target_field": "employee_id",
            "mapping_type": "direct",
        },
    ).json()
    client.post(f"/api/v1/mappings/{created['mapping_id']}/approve")
    preview = client.post(
        "/api/v1/migrations/preview",
        json={"source_id": source["source_id"], "mapping_id": created["mapping_id"]},
    ).json()
    await process_once(app.state.repository, "test-worker")

    export_resp = client.get(f"/api/v1/reports/{preview['job_id']}/export")
    assert export_resp.status_code == 200
    recon_resp = client.get(f"/api/v1/reports/{preview['job_id']}/reconciliation")
    assert recon_resp.status_code == 200

    actions = [e["action"] for e in app.state.repository.audits]
    assert "report.exported" in actions
    assert "report.reconciliation_accessed" in actions


def test_mapping_update_rejects_nonexistent_field(client: TestClient):
    source = upload(client)
    mapping = client.post(
        "/api/v1/mappings",
        json={
            "source_id": source["source_id"],
            "source_entity": "employee",
            "source_sheet": "Assignments",
            "source_field": "Employee ID",
            "target_entity": "employee",
            "target_field": "employee_id",
            "mapping_type": "unresolved",
        },
    ).json()
    bad_update = client.put(
        f"/api/v1/mappings/{mapping['mapping_id']}",
        json={"source_field": "Ghost_Field_99", "mapping_type": "direct"},
    )
    assert bad_update.status_code == 400
    assert bad_update.json()["error"]["code"] == "SOURCE_FIELD_NOT_FOUND"


@pytest.mark.asyncio
async def test_worker_cancellation_checkpoint(client: TestClient):
    from app.workers.worker import process_once

    source = upload(client)
    created = client.post(f"/api/v1/sources/{source['source_id']}/profile").json()
    job_id = created["job_id"]
    # Simulate cancellation requested while running
    await app.state.repository.request_cancel(job_id)
    # Worker processing should transition to CANCELLED
    await process_once(app.state.repository, "test-worker")
    job = await app.state.repository.get_job(job_id)
    assert job["status"] == "CANCELLED"


def test_list_jobs_and_sanitization(client: TestClient):
    source = upload(client)
    # Create profile job
    job1 = client.post(f"/api/v1/sources/{source['source_id']}/profile").json()
    # Create equus migration job
    job2 = client.post(
        f"/api/v1/sources/{source['source_id']}/equus-migration",
        json={"strategy": "EMPLOYEE_NUMBER"},
    ).json()

    # Query all jobs
    resp = client.get("/api/v1/jobs")
    assert resp.status_code == 200
    data = resp.json()
    assert data["total"] >= 2
    job_ids = [j["job_id"] for j in data["items"]]
    assert job1["job_id"] in job_ids
    assert job2["job_id"] in job_ids

    # Query filtered by kind
    filtered = client.get("/api/v1/jobs?kind=equus_migration").json()
    assert all(j["kind"] == "equus_migration" for j in filtered["items"])
    assert any(j["job_id"] == job2["job_id"] for j in filtered["items"])

    # Verify source_filename was populated
    target_job = next(j for j in filtered["items"] if j["job_id"] == job2["job_id"])
    assert target_job["source_filename"] == "client.xlsx"


@pytest.mark.asyncio
async def test_equus_migration_artifact_export_security(client: TestClient, tmp_path: Path):
    from app.core.config import settings

    source = upload(client)
    migration_job = client.post(
        f"/api/v1/sources/{source['source_id']}/equus-migration?strategy=TOPIA_ID_WITH_FALLBACK"
    ).json()
    job_id = migration_job["job_id"]

    # Simulate completed result pointing to authorized file
    authorized_dir = settings.report_dir / "equus"
    authorized_dir.mkdir(parents=True, exist_ok=True)
    real_import = authorized_dir / f"test_EQUUS_{job_id}.xlsx"
    real_import.write_bytes(b"PK\x03\x04test_content")

    repo = app.state.repository
    job = await repo.get_job(job_id)
    job["status"] = "COMPLETED"
    job["result"] = {
        "import_file": str(real_import),
        "report_file": str(authorized_dir / "missing.xlsx"),
    }
    await repo.put_job(job)

    # Reject unsupported artifact parameter
    bad_artifact = client.get(f"/api/v1/reports/{job_id}/export?artifact=malicious")
    assert bad_artifact.status_code == 400
    assert bad_artifact.json()["detail"]["code"] == "INVALID_ARTIFACT"

    # Valid export
    valid_resp = client.get(f"/api/v1/reports/{job_id}/export?artifact=import")
    assert valid_resp.status_code == 200
    assert valid_resp.content == b"PK\x03\x04test_content"

    # Missing file on disk
    missing_resp = client.get(f"/api/v1/reports/{job_id}/export?artifact=report")
    assert missing_resp.status_code == 404
    assert missing_resp.json()["detail"]["code"] == "ARTIFACT_NOT_FOUND"

    # Directory escape / path traversal rejection
    outside_file = tmp_path / "outside.xlsx"
    outside_file.write_bytes(b"secret")
    job["result"]["import_file"] = str(outside_file)
    await repo.put_job(job)

    escape_resp = client.get(f"/api/v1/reports/{job_id}/export?artifact=import")
    assert escape_resp.status_code == 403
    assert escape_resp.json()["detail"]["code"] == "ACCESS_DENIED"

    # Failed job: verify diagnostic header and filename prefix
    failed_import = authorized_dir / f"test_FAILED_{job_id}.xlsx"
    failed_import.write_bytes(b"PK\x03\x04failed_content")
    job["status"] = "FAILED"
    job["result"] = {
        "import_file": str(failed_import),
        "reconciliation": {"status": "FAILED", "is_valid": False, "errors": ["Missing Col 3"]},
    }
    await repo.put_job(job)

    failed_resp = client.get(f"/api/v1/reports/{job_id}/export?artifact=import")
    assert failed_resp.status_code == 200
    assert failed_resp.headers.get("X-Equus-Diagnostic-Only") == "true"
    assert "DIAGNOSTIC_ONLY_" in failed_resp.headers.get("Content-Disposition", "")


