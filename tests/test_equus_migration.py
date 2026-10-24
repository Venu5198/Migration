"""Tests for the Topia to Equus Migration Engine."""

from pathlib import Path

import openpyxl
import pytest

from app.repositories import MemoryRepository
from app.schemas import JobStatus
from app.services.equus import (
    DEFAULT_TEMPLATE_PATH,
    parse_three_row_workbook,
    reconcile_equus_transformation,
    run_equus_migration_pipeline,
    transform_topia_to_equus,
)
from app.services.equus.mapping_config import (
    STRATEGY_EMPLOYEE_NUMBER,
    STRATEGY_TOPIA_ID_WITH_FALLBACK,
    clean_not_provided,
    convert_country_to_iso3,
    convert_dropdown,
    convert_us_date,
    normalize_identifier,
    split_multiline_address,
)
from app.services.equus.transformer import TransformationResult
from app.workers.worker import process_once


@pytest.fixture
def sample_template(tmp_path: Path) -> Path:
    """Create a minimal 3-row Equus template for fast unit testing."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    # Row 1: Instructions
    ws.append(
        [
            "Instructions --->>>",
            None,
            None,
            None,
            None,
            "Match dropdown",
            "US Date Format",
            "Match Country Code",
            "Match dropdown",
            None,
            None,
            None,
        ]
    )
    # Row 2: Field Name / Label
    ws.append(
        [
            "Field Name --->>>",
            None,
            "Talent Link Id",
            None,
            "Topia Id",
            "Gender",
            "Birth Date",
            "Home Country",
            "Location Type",
            "Address Line 1",
            "Address Line 2",
            "Personal Mobile",
        ]
    )
    # Row 3: Technical Header
    ws.append(
        [
            "A",
            "Assignment Command",
            "Employee Number",
            "Assignment Number",
            "Assignment Identifier",
            "EMPLOYEE!GENDER",
            "EMPLOYEE!BIRTH_DATE",
            "ASSIGNMENT!FROM_COUNTRY",
            "ASSIGNMENT_MAILING_ADDRESS!LOCATION_TYPE",
            "ASSIGNMENT_MAILING_ADDRESS!LINE_1",
            "ASSIGNMENT_MAILING_ADDRESS!LINE_2",
            "ASSIGNMENT_EMPLOYEE_CONTACT!NUMBER_ADDRESS",
        ]
    )
    # Sample row 4
    ws.append(
        [
            "INSERT",
            "INSERT",
            "E001",
            "NEW",
            "T001",
            "Male",
            "01/01/1980",
            "USA",
            "Destination Office",
            "1 Office Park",
            None,
            "123456",
        ]
    )
    p = tmp_path / "test_template.xlsx"
    wb.save(p)
    return p


@pytest.fixture
def synthetic_source(tmp_path: Path) -> Path:
    """Create a synthetic Topia 4-row header source workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    # Row 1: Rules
    ws.append(
        [
            None,
            None,
            None,
            None,
            None,
            "Match dropdown",
            "US Date Format",
            "Match Country Code",
            None,
            None,
            None,
        ]
    )
    # Row 2: Technical Field
    ws.append(
        [
            None,
            None,
            "EMPLOYEE!SYSTEM_NUMBER",
            None,
            "ASSIGNMENT_CF!CF_TEXT029",
            "EMPLOYEE!GENDER",
            "EMPLOYEE!BIRTH_DATE",
            "ASSIGNMENT!FROM_COUNTRY",
            "ASSIGNMENT_MAILING_ADDRESS!LOCATION_TYPE",
            "ASSIGNMENT_EMPLOYEE_CONTACT!NUMBER_ADDRESS",
            "ASSIGNMENT_MAILING_ADDRESS!LINE_1",
        ]
    )
    # Row 3: Label
    ws.append(
        [
            None,
            None,
            "Talent Link ID",
            None,
            "Topia ID",
            "Gender",
            "Date of Birth",
            "Home Country",
            "Work Host Address",
            "Personal Phone",
            "Personal Home Address with Effective date",
        ]
    )
    # Row 4: Display Header
    ws.append(
        [
            "Full Name",
            "First Name",
            "Employee ID",
            None,
            "Assignment ID",
            "Gender",
            "Birth Date",
            "Home Country",
            "Work Host Address",
            "Phone Number",
            "Personal Home Address",
        ]
    )
    # Row 5: Data 1 (Clean: 2 addresses for John Smith -> produces 1 INSERT + 1 UPDATE)
    ws.append(
        [
            "Smith, John",
            "John",
            "1001",
            None,
            "A1001",
            "Male",
            "1985-06-15",
            "United States",
            "1 Office Park",
            "555-0100",
            "12 Big Street\nDown the Road",
        ]
    )
    # Row 6: Data 2 (Jane Doe: 1 row with 'Not Provided' and ISO-3)
    ws.append(
        [
            "Doe, Jane",
            "Jane",
            "1002",
            None,
            "A1002",
            "Female",
            "07/20/1990",
            "India",
            "Not Provided",
            None,
            None,
        ]
    )
    p = tmp_path / "test_source.xlsx"
    wb.save(p)
    return p


def test_clean_not_provided():
    assert clean_not_provided("Not Provided") is None
    assert clean_not_provided("N/A") is None
    assert clean_not_provided("na") is None
    assert clean_not_provided("-") is None
    assert clean_not_provided("") is None
    assert clean_not_provided("Valid Value") == "Valid Value"


def test_date_conversion():
    d1, err1 = convert_us_date("1985-06-15")
    assert err1 is None
    assert d1 == "06/15/1985"

    d2, err2 = convert_us_date("15/06/1985")
    assert err2 is None
    assert d2 == "06/15/1985"

    d3, err3 = convert_us_date("invalid_date")
    assert d3 is None
    assert "could not be parsed" in err3


def test_country_to_iso3():
    c1, err1 = convert_country_to_iso3("India")
    assert err1 is None
    assert c1 == "IND"

    c2, err2 = convert_country_to_iso3("United States")
    assert err2 is None
    assert c2 == "USA"

    c3, err3 = convert_country_to_iso3("JP")
    assert err3 is None
    assert c3 == "JPN"

    c4, err4 = convert_country_to_iso3("Atlantis")
    assert c4 is None
    assert "UNKNOWN_COUNTRY_CODE" in err4


def test_dropdown_conversion():
    v1, err1 = convert_dropdown("EMPLOYMENT_STATUS", "Employed")
    assert err1 is None
    assert v1 == "Active"

    v2, err2 = convert_dropdown("GENDER", "female")
    assert err2 is None
    assert v2 == "Female"

    v3, err3 = convert_dropdown("YES_NO", "true")
    assert err3 is None
    assert v3 == "Yes"

    v4, err4 = convert_dropdown("GENDER", "Alien")
    assert v4 is None
    assert "UNKNOWN_DROPDOWN_VALUE" in err4


def test_parse_three_row_workbook(synthetic_source: Path):
    parsed = parse_three_row_workbook(synthetic_source)
    assert parsed.header_rows_count == 4
    assert len(parsed.data_rows) == 2
    assert parsed.total_rows_evaluated == 2
    assert len(parsed.columns) == 11


def test_transform_topia_to_equus(synthetic_source: Path, sample_template: Path):
    parsed_src = parse_three_row_workbook(synthetic_source)
    res = transform_topia_to_equus(parsed_src, sample_template)

    assert res.total_evaluated_rows == 2
    assert len(res.insert_records) == 2
    # Row 1 has 2 addresses, producing 1 UPDATE record
    assert len(res.update_records) >= 1
    assert len(res.all_output_records) == len(res.insert_records) + len(res.update_records)

    # Check first INSERT row (under Option B / STRATEGY_EMPLOYEE_NUMBER default)
    r1 = res.insert_records[0]
    assert r1[1] == "INSERT"
    assert r1[2] == "INSERT"
    assert r1[3] == "1001"
    assert r1[4] == "NEW"
    assert r1[5] == "1001"  # Option B: Employee Number is in Col 5
    assert r1[6] == "Male"
    assert r1[7] == "06/15/1985"
    assert r1[8] == "USA"
    # Child 0 nested on INSERT row
    assert r1.get(9) == "Destination Office"
    assert r1.get(10) == "1 Office Park"
    assert r1.get(12) == "555-0100"

    # Check UPDATE row
    u1 = res.update_records[0]
    assert u1[1] == "IGNORE"
    assert u1[2] == "UPDATE"
    assert u1.get(3) is None  # Strictly blank on UPDATE
    assert u1.get(4) is None  # Strictly blank on UPDATE
    assert u1[5] == "1001"  # References parent Assignment Identifier
    assert u1.get(9) == "Departure Residence"
    assert u1.get(10) == "12 Big Street"
    assert u1.get(11) == "Down the Road"


def test_reconciliation_detects_duplicates(synthetic_source: Path, sample_template: Path):
    parsed_src = parse_three_row_workbook(synthetic_source)
    res = transform_topia_to_equus(parsed_src, sample_template)

    # Inject duplicate assignment ID
    res.insert_records[1][5] = res.insert_records[0][5]
    recon = reconcile_equus_transformation(res)
    assert recon["is_valid"] is False
    assert recon["status"] == "FAILED"
    assert any("Duplicate Assignment Identifier" in err for err in recon["errors"])


def test_full_pipeline_artifacts_and_report(
    synthetic_source: Path, sample_template: Path, tmp_path: Path
):
    out_dir = tmp_path / "artifacts"
    res = run_equus_migration_pipeline(
        source_path=synthetic_source,
        template_path=sample_template,
        output_dir=out_dir,
        source_name="smoke_test",
    )
    assert res["status"] in ("VERIFIED_PASS", "COMPLETED_WITH_WARNINGS")
    assert Path(res["import_file"]).exists()
    assert Path(res["report_file"]).exists()

    # Verify generated Equus import file
    wb_import = openpyxl.load_workbook(res["import_file"])
    ws_import = wb_import.active
    assert ws_import.cell(4, 1).value == "INSERT"
    assert ws_import.cell(4, 3).value == "1001"
    wb_import.close()

    # Verify generated multi-tab report
    wb_report = openpyxl.load_workbook(res["report_file"])
    assert "Executive Summary" in wb_report.sheetnames
    assert "Column Mappings" in wb_report.sheetnames
    assert "Translations Applied" in wb_report.sheetnames
    assert "Warnings & Issues" in wb_report.sheetnames
    wb_report.close()


@pytest.mark.asyncio
async def test_worker_processes_equus_migration_job(synthetic_source: Path, tmp_path: Path):
    repo = MemoryRepository()
    source_id = "src-equus-1"
    await repo.put_source(
        {
            "source_id": source_id,
            "original_filename": "synthetic_source.xlsx",
            "storage_path": str(synthetic_source),
            "sha256": "test-sha",
            "sheets": [],
        }
    )

    from datetime import UTC, datetime

    job_id = "job-equus-1"
    now = datetime.now(UTC)
    await repo.put_job(
        {
            "job_id": job_id,
            "kind": "equus_migration",
            "source_id": source_id,
            "status": JobStatus.PENDING,
            "progress": 0,
            "attempts": 0,
            "created_at": now,
            "updated_at": now,
            "cancellation_requested": False,
        }
    )

    processed = await process_once(repo, "test-worker")
    assert processed is True

    finished_job = await repo.get_job(job_id)
    assert finished_job["status"] in (JobStatus.COMPLETED, JobStatus.COMPLETED_WITH_ERRORS)
    assert "import_file" in finished_job["result"]
    assert "report_file" in finished_job["result"]
    assert finished_job["result"]["insert_records"] == 2


# ===========================================================================
# Phase 2 Regression Tests: Verified Migration Business Rules
# ===========================================================================


def test_regression_insert_with_valid_assignment_identifier():
    """1. INSERT with a valid Assignment Identifier passes reconciliation."""
    res = TransformationResult(
        insert_records=[{1: "INSERT", 2: "INSERT", 3: "60936", 4: "NEW", 5: "19338"}],
        update_records=[],
        all_output_records=[{1: "INSERT", 2: "INSERT", 3: "60936", 4: "NEW", 5: "19338"}],
        total_evaluated_rows=1,
    )
    recon = reconcile_equus_transformation(res)
    assert recon["is_valid"] is True
    assert recon["status"] == "VERIFIED_PASS"
    assert len(recon["errors"]) == 0


def test_regression_update_with_blank_employee_number_and_valid_parent():
    """2. UPDATE with a blank Employee Number and a valid parent identifier passes."""
    res = TransformationResult(
        insert_records=[{1: "INSERT", 2: "INSERT", 3: "60936", 4: "NEW", 5: "19338"}],
        update_records=[
            {
                1: "IGNORE",
                2: "UPDATE",
                3: None,
                4: None,
                5: "19338",
                9: "Destination Residence",
            }
        ],
        all_output_records=[
            {1: "INSERT", 2: "INSERT", 3: "60936", 4: "NEW", 5: "19338"},
            {
                1: "IGNORE",
                2: "UPDATE",
                3: None,
                4: None,
                5: "19338",
                9: "Destination Residence",
            },
        ],
        total_evaluated_rows=1,
    )
    recon = reconcile_equus_transformation(res)
    assert recon["is_valid"] is True
    assert recon["status"] == "VERIFIED_PASS"
    assert len(recon["errors"]) == 0


def test_regression_missing_insert_identifier():
    """3. Missing INSERT identifier must fail validation with row-type and field detail."""
    res = TransformationResult(
        insert_records=[{1: "INSERT", 2: "INSERT", 3: "60936", 4: "NEW", 5: None}],
        update_records=[],
        all_output_records=[{1: "INSERT", 2: "INSERT", 3: "60936", 4: "NEW", 5: None}],
        total_evaluated_rows=1,
    )
    recon = reconcile_equus_transformation(res)
    assert recon["is_valid"] is False
    assert recon["status"] == "FAILED"
    assert any("missing required Assignment Identifier (Col 5)" in err for err in recon["errors"])


def test_regression_update_referencing_unknown_parent():
    """4. UPDATE referencing an unknown parent must fail validation."""
    res = TransformationResult(
        insert_records=[{1: "INSERT", 2: "INSERT", 3: "60936", 4: "NEW", 5: "19338"}],
        update_records=[
            {
                1: "IGNORE",
                2: "UPDATE",
                3: None,
                4: None,
                5: "UNKNOWN_ID",
                9: "Destination Residence",
            }
        ],
        all_output_records=[
            {1: "INSERT", 2: "INSERT", 3: "60936", 4: "NEW", 5: "19338"},
            {
                1: "IGNORE",
                2: "UPDATE",
                3: None,
                4: None,
                5: "UNKNOWN_ID",
                9: "Destination Residence",
            },
        ],
        total_evaluated_rows=1,
    )
    recon = reconcile_equus_transformation(res)
    assert recon["is_valid"] is False
    assert recon["status"] == "FAILED"
    assert any(
        "references unparented Assignment Identifier: 'UNKNOWN_ID'" in err
        for err in recon["errors"]
    )


def test_regression_multiple_addresses_contacts_packing():
    """5. Multiple addresses, contacts, and multiline address splitting."""
    l1, l2, l3 = split_multiline_address("12 Big Street\nDown the Road\nApartment 4B")
    assert l1 == "12 Big Street"
    assert l2 == "Down the Road"
    assert l3 == "Apartment 4B"

    # Single line address
    s1, s2, s3 = split_multiline_address("1 Office Park")
    assert s1 == "1 Office Park"
    assert s2 is None
    assert s3 is None


def test_regression_numeric_identifier_normalization():
    """6. Numeric identifier normalization preserves integers without trailing .0."""
    assert normalize_identifier(60936.0) == "60936"
    assert normalize_identifier("19338.0") == "19338"
    assert normalize_identifier(10500545) == "10500545"
    assert normalize_identifier("A1001") == "A1001"
    assert normalize_identifier("Not Provided") is None
    assert normalize_identifier(None) is None


def test_regression_client_expected_workbook_structure():
    """7. The client expected workbook's INSERT/UPDATE structure."""
    client_records = [
        # Row 4 (Data 1 / INSERT): Employee 60936 with primary address, contact
        {
            1: "INSERT",
            2: "INSERT",
            3: "60936",
            4: "NEW",
            5: "60936",
            9: "Destination Office",
            12: "dm@null.null",
        },
        # Row 5 (Data 2 / UPDATE): Child update, Col 3 and 4 strictly None
        {
            1: "IGNORE",
            2: "UPDATE",
            3: None,
            4: None,
            5: "60936",
            9: "Departure Residence",
            12: "1234567",
        },
        # Row 6 (Data 3 / UPDATE): Address 3
        {1: "IGNORE", 2: "UPDATE", 3: None, 4: None, 5: "60936", 9: "Destination Residence"},
        # Row 7 (Data 4 / UPDATE): Address 4
        {1: "IGNORE", 2: "UPDATE", 3: None, 4: None, 5: "60936", 9: "Departure Office"},
    ]
    res = TransformationResult(
        insert_records=[client_records[0]],
        update_records=client_records[1:],
        all_output_records=client_records,
        total_evaluated_rows=1,
    )
    recon = reconcile_equus_transformation(res)
    assert recon["is_valid"] is True
    assert recon["status"] == "VERIFIED_PASS"
    assert recon["insert_records_count"] == 1
    assert recon["update_records_count"] == 3
    assert recon["expansion_ratio"] == 4.0


def test_regression_mapping_behavior_for_unresolved_identifier_choices(
    synthetic_source: Path, sample_template: Path
):
    """8. Mapping behavior for unresolved identifier choices (strategies)."""
    parsed_src = parse_three_row_workbook(synthetic_source)

    # Strategy A: Reference strategy (TOPIA_ID_WITH_FALLBACK)
    res_a = transform_topia_to_equus(
        parsed_src, sample_template, strategy=STRATEGY_TOPIA_ID_WITH_FALLBACK
    )
    assert res_a.insert_records[0][5] == "A1001"  # Uses Topia ID

    # Strategy B: Client Test File strategy (EMPLOYEE_NUMBER)
    res_b = transform_topia_to_equus(parsed_src, sample_template, strategy=STRATEGY_EMPLOYEE_NUMBER)
    assert res_b.insert_records[0][5] == "1001"  # Uses Employee Number in Col 5


def test_reconciliation_count_consistency_and_semantics():
    """Verify reconciliation count aliases and accepted/rejected semantics agree."""
    # 1. Valid result: all records accepted
    valid_res = TransformationResult(
        insert_records=[{1: "INSERT", 2: "INSERT", 3: "E101", 4: "NEW", 5: "E101"}],
        update_records=[],
        all_output_records=[{1: "INSERT", 2: "INSERT", 3: "E101", 4: "NEW", 5: "E101"}],
        total_evaluated_rows=1,
    )
    recon_valid = reconcile_equus_transformation(valid_res)
    assert recon_valid["is_valid"] is True
    assert recon_valid["status"] == "VERIFIED_PASS"
    assert recon_valid["source_rows"] == recon_valid["source_rows_evaluated"] == 1
    assert recon_valid["output_records"] == recon_valid["total_output_records"] == 1
    assert recon_valid["attempted_records"] == 1
    assert recon_valid["accepted_records"] == 1
    assert recon_valid["rejected_records"] == 0

    # 2. Invalid result (missing required fields): all records rejected from import
    invalid_res = TransformationResult(
        insert_records=[{1: "INSERT", 2: "INSERT", 3: None, 4: "NEW", 5: None}],
        update_records=[],
        all_output_records=[{1: "INSERT", 2: "INSERT", 3: None, 4: "NEW", 5: None}],
        total_evaluated_rows=1,
    )
    recon_invalid = reconcile_equus_transformation(invalid_res)
    assert recon_invalid["is_valid"] is False
    assert recon_invalid["status"] == "FAILED"
    assert recon_invalid["source_rows"] == recon_invalid["source_rows_evaluated"] == 1
    assert recon_invalid["output_records"] == recon_invalid["total_output_records"] == 1
    assert recon_invalid["attempted_records"] == 1
    assert recon_invalid["accepted_records"] == 0
    assert recon_invalid["rejected_records"] == 1
    assert recon_invalid["validation_errors_count"] == 2


def test_negative_fixture_report_and_reconciliation_status_match(tmp_path: Path):
    """Negative fixture generates FAILED reconciliation and report Tab 1 + Tab 4 reflect errors."""
    import openpyxl

    neg_fixture = Path("topia_equus_test_workbooks/04_negative_missing_employee_id.xlsx")
    if not neg_fixture.exists():
        pytest.skip("Negative fixture not present")

    res = run_equus_migration_pipeline(
        source_path=neg_fixture,
        template_path=DEFAULT_TEMPLATE_PATH,
        output_dir=tmp_path,
        overwrite=True,
    )

    assert res["status"] == "FAILED"
    assert res["reconciliation"]["is_valid"] is False
    assert res["reconciliation"]["validation_errors_count"] == 2
    assert res["accepted_records"] == 0
    assert res["rejected_records"] == res["total_output_records"]

    # Verify report Excel contents
    wb = openpyxl.load_workbook(res["report_file"])
    ws_summary = wb["Executive Summary"]
    summary_dict = {
        ws_summary.cell(row=r, column=1).value: ws_summary.cell(row=r, column=2).value
        for r in range(1, ws_summary.max_row + 1)
    }
    assert summary_dict["Reconciliation Status"] == "FAILED"
    assert summary_dict["Validation Errors"] == 2
    assert summary_dict["Accepted Import Records"] == 0
    assert summary_dict["Rejected / Diagnostic Records"] == res["total_output_records"]

    # Verify Tab 4 contains the validation errors
    ws_warn = wb["Warnings & Issues"]
    rows = [
        [ws_warn.cell(row=r, column=c).value for c in range(1, 6)]
        for r in range(2, ws_warn.max_row + 1)
    ]
    assert len(rows) >= 2
    codes = [r[0] for r in rows]
    assert "REQUIRED_FIELD_ERROR" in codes
    messages = " ".join([str(r[4]) for r in rows])
    assert "Employee Number" in messages
    assert "Assignment Identifier" in messages


def test_valid_fixture_report_and_reconciliation_status_match(tmp_path: Path):
    """Valid basic fixture produces VERIFIED_PASS in both reconciliation and mapping report."""
    import openpyxl

    valid_fixture = Path("topia_equus_test_workbooks/01_valid_basic.xlsx")
    if not valid_fixture.exists():
        pytest.skip("Valid fixture not present")

    res = run_equus_migration_pipeline(
        source_path=valid_fixture,
        template_path=DEFAULT_TEMPLATE_PATH,
        output_dir=tmp_path,
        overwrite=True,
    )

    assert res["status"] == "VERIFIED_PASS"
    assert res["reconciliation"]["is_valid"] is True
    assert res["accepted_records"] == res["total_output_records"]
    assert res["rejected_records"] == 0

    wb = openpyxl.load_workbook(res["report_file"])
    ws_summary = wb["Executive Summary"]
    summary_dict = {
        ws_summary.cell(row=r, column=1).value: ws_summary.cell(row=r, column=2).value
        for r in range(1, ws_summary.max_row + 1)
    }
    assert summary_dict["Reconciliation Status"] == "VERIFIED_PASS"
    assert summary_dict["Validation Errors"] == 0
    assert summary_dict["Accepted Import Records"] == res["total_output_records"]


def test_destination_writes_blocked_on_reconciliation_failure(tmp_path: Path):
    """Destination execution is blocked if reconciliation fails, even if demo writes requested."""
    neg_fixture = Path("topia_equus_test_workbooks/04_negative_missing_employee_id.xlsx")
    if not neg_fixture.exists():
        pytest.skip("Negative fixture not present")

    res = run_equus_migration_pipeline(
        source_path=neg_fixture,
        template_path=DEFAULT_TEMPLATE_PATH,
        output_dir=tmp_path,
        demo_writes_enabled=True,
        overwrite=True,
    )

    assert res["status"] == "FAILED"
    # Destination writes must NOT have executed
    exec_summary = res["execution_summary"]
    assert exec_summary["status"] == "BLOCKED_RECONCILIATION_FAILED"
    assert exec_summary["attempted"] == 0


def test_strategy_consistency_in_pipeline_and_report(tmp_path: Path):
    """Selected strategy is propagated to pipeline result and mapping report."""
    import openpyxl

    valid_fixture = Path("topia_equus_test_workbooks/01_valid_basic.xlsx")
    if not valid_fixture.exists():
        pytest.skip("Valid fixture not present")

    # Test STRATEGY_TOPIA_ID_WITH_FALLBACK
    res = run_equus_migration_pipeline(
        source_path=valid_fixture,
        template_path=DEFAULT_TEMPLATE_PATH,
        output_dir=tmp_path,
        strategy=STRATEGY_TOPIA_ID_WITH_FALLBACK,
        overwrite=True,
    )
    assert res["strategy"] == STRATEGY_TOPIA_ID_WITH_FALLBACK
    assert res["reconciliation"]["assignment_id_strategy"] == STRATEGY_TOPIA_ID_WITH_FALLBACK

    wb = openpyxl.load_workbook(res["report_file"])
    ws_summary = wb["Executive Summary"]
    summary_dict = {
        ws_summary.cell(row=r, column=1).value: ws_summary.cell(row=r, column=2).value
        for r in range(1, ws_summary.max_row + 1)
    }
    assert summary_dict["Assignment Identifier Strategy"] == STRATEGY_TOPIA_ID_WITH_FALLBACK
